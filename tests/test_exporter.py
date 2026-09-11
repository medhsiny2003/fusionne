"""
Tests unitaires pour le module d'exportation Excel.
"""

import io
import os
import openpyxl
import pandas as pd
import pytest
from core.exporter import export_to_excel, generate_export_filename, MX_STYLES
from core.file_reader import EXPECTED_COLUMNS

def test_export_to_excel_buffer():
    """Vérifie la génération du buffer binaire Excel stylisé."""
    df = pd.DataFrame([
        {
            "Prénom": "Alice",
            "Nom": "Vasseur",
            "Poste Actuel": "CEO",
            "Entreprise": "Innovate",
            "Email Proposé": "alice@innovate.com",
            "Email Alternatif 1": "",
            "Email Alternatif 2": "",
            "Score de Confiance (%)": 98,
            "Statut MX": "Validé",
            "Serveur MX Actif": "Oui",
            "Lien Profil LinkedIn": "linkedin.com/in/alice",
            "Date d'Extraction": "2026-09-01"
        },
        {
            "Prénom": "Bob",
            "Nom": "Durand",
            "Poste Actuel": "CTO",
            "Entreprise": "TechSoft",
            "Email Proposé": "bob@techsoft.com",
            "Email Alternatif 1": "",
            "Email Alternatif 2": "",
            "Score de Confiance (%)": 45,
            "Statut MX": "Invalide",
            "Serveur MX Actif": "Non",
            "Lien Profil LinkedIn": "linkedin.com/in/bob",
            "Date d'Extraction": "2026-09-01"
        },
    ])

    buffer = export_to_excel(df)
    assert isinstance(buffer, io.BytesIO)
    buffer.seek(0)

    # Recharger avec openpyxl pour vérifier les styles
    wb = openpyxl.load_workbook(buffer)
    ws = wb.active
    assert ws.title == "Contacts Fusionnés"
    assert ws.max_row == 3 # 1 header + 2 data
    assert ws.max_column == len(EXPECTED_COLUMNS)

    # Vérification des couleurs d'en-tête
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.fill.start_color.rgb == "000A66C2" or header_cell.fill.start_color.rgb == "0A66C2"
    assert header_cell.font.bold is True

    # Vérification de l'auto-filtre
    assert ws.auto_filter.ref is not None

def test_export_to_file(tmp_path):
    """Vérifie l'écriture de l'exportation sur le disque."""
    df = pd.DataFrame([{"Prénom": "Test", "Nom": "User", "Statut MX": "Validé"}])
    export_file = tmp_path / "test_out.xlsx"
    
    export_to_excel(df, output_path=str(export_file))
    assert os.path.exists(export_file)
    assert os.path.getsize(export_file) > 0

def test_generate_export_filename():
    """Vérifie le format du nom de fichier généré."""
    filename = generate_export_filename()
    assert filename.startswith("contacts_fusionnes_")
    assert filename.endswith(".xlsx")
